"""KiCad schematic -> BOM workbook (xlsx + pdf) in the RETI 'BILL OF MATERIAL LIST' layout.

raw.csv comes from kicad-cli, one row per reference:
  kicad-cli sch export bom --ref-range-delimiter "" -o raw.csv <root.kicad_sch> --fields
    "Reference,Value,Footprint,Manufacturer,Manufacturer_Part_Number,MPN,Description,Voltage,Dielectric,
     Tolerance,Rating,${DNP},${EXCLUDE_FROM_BOM}"
  --labels "Ref,Value,Footprint,Mfr,MPN,MPN2,Desc,Voltage,Dielectric,Tol,Rating,DNP,ExBOM"
(MVP-AMP raw.csv was exported without MPN2 and with an extra Irated/AsmNote column - both shapes work.)

usage: build_bom.py raw.csv out.xlsx "<commit label>" --sch <root.kicad_sch> --model "<RETI MODEL>"
                    --source "<repo / file>" [--supp bom_sources.json]
"""
import argparse, csv, json, re, sys, os, math, datetime
from collections import OrderedDict, defaultdict
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.worksheet.page import PageMargins

sys.stdout.reconfigure(encoding='utf-8')
ap = argparse.ArgumentParser()
ap.add_argument('raw')
ap.add_argument('out')
ap.add_argument('commit')
ap.add_argument('--sch', required=True, help='root .kicad_sch; its folder holds the project *.pretty libraries')
ap.add_argument('--model', required=True, help='RETI MODEL header value, e.g. "MVP-AMP (usd_amp)"')
ap.add_argument('--source', required=True, help='source line for the review sheet')
ap.add_argument('--supp', help='bom_sources.json with supplementary specs from an earlier reviewed BOM')
a = ap.parse_args()
RAW, OUT, COMMIT = a.raw, a.out, a.commit
SCH_DIR = os.path.dirname(os.path.abspath(a.sch))
KICAD_FP = r'C:\Program Files\KiCad\10.0\share\kicad\footprints'

rows = list(csv.DictReader(open(RAW, encoding='utf-8')))
SUPP = {}
if a.supp:
    SUPP = {p['Reference']: p for p in json.load(open(a.supp, encoding='utf-8'))['parts']}

# ---------------------------------------------------------------- footprints
_fp_cache = {}
def fp_attr_of(fp_full):
    """'(attr smd)' / 'through_hole' of a 'Lib:Name' footprint: project lib first, then the KiCad stock libs."""
    if fp_full not in _fp_cache:
        lib, name = fp_full.split(':', 1)
        _fp_cache[fp_full] = None
        for d in (os.path.join(SCH_DIR, lib + '.pretty'), os.path.join(KICAD_FP, lib + '.pretty')):
            f = os.path.join(d, name + '.kicad_mod')
            if os.path.exists(f):
                m = re.search(r'\(attr ([a-z_ ]+)\)', open(f, encoding='utf-8').read())
                _fp_cache[fp_full] = m.group(1) if m else ''
                break
    return _fp_cache[fp_full]

# ---------------------------------------------------------------- symbols left out of the BOM (in_bom no)
def sheet_files(root):
    seen, todo = [], [os.path.abspath(root)]
    while todo:
        f = todo.pop()
        if f in seen or not os.path.exists(f):
            continue
        seen.append(f)
        t = open(f, encoding='utf-8').read()
        todo += [os.path.join(os.path.dirname(f), s) for s in re.findall(r'\(property "Sheetfile" "([^"]+)"', t)]
    return seen

excluded = defaultdict(list)            # lib symbol name -> refs
for f in sheet_files(a.sch):
    t = open(f, encoding='utf-8').read()
    for m in re.finditer(r'\n\t\(symbol\n\t\t\(lib_id "([^"]+)"(.*?)\n\t\)', t, re.S):
        body = m.group(2)
        if '(in_bom no)' in body:
            ref = re.search(r'\(property "Reference" "([^"]+)"', body).group(1)
            if not ref.startswith('#'):
                excluded[m.group(1).split(':')[-1]].append(ref)

IMP2MET = {'0402': '1005', '0603': '1608', '0805': '2012', '1206': '3216',
           '1210': '3225', '1812': '4532', '2512': '6332'}

def size_of(fp):
    m = re.search(r'(\d{4})Metric', fp)
    if m:
        return m.group(1)
    m = re.search(r'(0402|0603|0805|1206|1210|1812|2512)', fp)
    return IMP2MET[m.group(1)] if m else ''

PKG = [  # footprint substring -> PACKAGE label
    ('CP_Radial_D10.0mm', 'ECAP-D10(P5.0)'), ('D_SMA', 'SMA(DO-214AC)'), ('D_SMB', 'SMB(DO-214AA)'),
    ('D_SMF', 'SMF(SOD-123FL)'), ('D_SOD-123', 'SOD-123'), ('SOT-23-5', 'SOT23-5'),
    ('TSOT-23-6', 'TSOT23-6'), ('SOT-23-6', 'SOT23-6'), ('SOT-23', 'SOT-23'),
    ('SOT-323', 'SOT-323'), ('SOT-353', 'SOT-353(SC70-5)'), ('WSON-6', 'WSON-6(2x2)'),
    ('EPC8010_LGA6', 'LGA6(2.1x0.85)'), ('XGD10603', '0603(1608)'),
    ('Fuse_Bourns_SF-1206F', '1206(3216)'), ('MF-MSMF_Style2_1812', '1812(4532)'),
    ('NANO2', 'NANO2(451/453)'), ('LED_0603', 'LED1608'), ('XAL6060', 'XAL6060'),
    ('IHLP-5050', 'IHLP-5050'), ('IHLP-4040', 'IHLP-4040'), ('WE-LQ-1812', '1812(4532)'),
    ('Coilcraft_1812FS', '1812(4532)'), ('G6K-2F-Y', 'G6K-2F-Y(SMD)'),
    ('SO-8FL', 'SO-8FL(5x6)'), ('PowerPAK_1212-8', 'PowerPAK1212-8'),
    ('VSON-8_3.3x3.3', 'VSON-8(3.3x3.3)'), ('RM6', 'RM6(6P THT)'), ('CST2010', 'CST2010'),
    ('NTCS0603', '0603(1608)'), ('HSOP-8-1EP', 'HSOIC-8-EP(DDA)'), ('HTSOP-8-1EP', 'HSOIC-8-EP(DDA)'),
    ('DRC0010J', 'VSON-10(DRC 3x3)'), ('RPM0026A', 'VQFN-HR-26(3.5x4)'),
    ('SOIC-16', 'SOIC-16'), ('SOIC-18W', 'SOIC-18W'), ('SOIC-14', 'SOIC-14'),
    ('TSSOP-14', 'TSSOP-14'), ('DSBGA-6', 'DSBGA-6(YFF)'), ('VQFN-24', 'VQFN-24(4x4)'),
    ('MOD_BoT_nLE523D', 'DIP 2x4 P2.54'), ('Nucleo144', 'Morpho CN11/CN12'),
    ('Buzzer_12x9.5', 'BUZZER-12x9.5'), ('B2B-XH', 'XH-2P'), ('B4B-XH', 'XH-4P'),
    ('B2P-VH', 'VH-2P'), ('B3P-VH', 'VH-3P'), ('PinHeader_1x07', 'HEADER_7P_2.54'),
    ('SMA_ACERFCOMM', 'SMA-F(R/A)'), ('PJ-063BH', 'DCJACK-PJ063BH'), ('B3S-1000', 'B3S-1000(SMD)'),
    ('TestPoint_Pad_D1.5', 'TP-PAD-D1.5'),
    # MVP-PPG
    ('CablePad_D1.2', 'PAD-D1.2'), ('TestPad_D0.8', 'PAD-D0.8'), ('TVS_SOD882', 'SOD-882(1006)'),
    ('TI_DPY0002A', 'X1SON-2 DPY(1006)'), ('MAXM86161_OLGA-14', 'OLGA-14(2.9x4.3)'),
    ('UFQFPN-20', 'UFQFPN-20(3x3)'), ('TSOT23-6_MaxLinear', 'TSOT23-6'), ('DFN-8-1EP_2x3mm', 'DFN-8-EP(2x3)'),
]
def pkg_of(fp):
    for k, v in PKG:
        if k in fp:
            return v
    return fp

# ---------------------------------------------------------------- helpers
def num(s):
    s = ('%g' % float(s))
    return s

def cap_fmt(pf):
    if pf < 1000:
        return num(pf) + 'pF'
    if pf < 100000:
        return num(pf / 1000) + 'nF'
    return num(pf / 1e6) + 'uF'

CAP_MULT = {'pF': 1, 'nF': 1e3, 'uF': 1e6}

def refkey(ref):
    m = re.match(r'([A-Z]+)(\d+)', ref)
    return (m.group(1), int(m.group(2)))

def mlcc_strings(pf, v, diel, tol, size):
    part = cap_fmt(pf) + (f'/{v}V' if v else '') + ('' if size == '1608' else f'/{size}')
    sd = '/'.join(x for x in ['MLCC', diel, (v + 'V') if v else '', tol if tol and tol != '5%' else '', size] if x)
    return part, sd

def split_vendor(mfr):
    """'Nexperia BAT54GWJ (alt. Diodes 1N5819HW-7-F)' -> ('Nexperia', 'BAT54GWJ', 'Diodes 1N5819HW-7-F')"""
    alt = ''
    m = re.search(r'\((alt\.\s*)?([^)]*)\)?$', mfr)
    if m:
        inner = m.group(2).strip()
        if m.group(1):
            alt = '대체: ' + inner
            mfr = mfr[:m.start()].strip()
        elif re.search(r'\d|%', inner):      # spec remark, e.g. 'FHT for 1 % R25'
            alt = inner
            mfr = mfr[:m.start()].strip()
        else:                                 # alias, e.g. 'Same Sky (CUI Devices)'
            mfr = mfr[:m.start()].strip() + ' (' + inner + ')'
    if mfr.lower().startswith('generic'):
        mfr = 'Any'
    pn = ''
    toks = mfr.split()
    if len(toks) >= 2 and re.search(r'\d', toks[-1]) and len(toks[-1]) >= 6 and toks[-1].upper() == toks[-1]:
        pn = toks[-1]
        mfr = ' '.join(toks[:-1])
    return mfr.strip(), pn, alt

def pn_from_desc(val, desc):
    tok = desc.split(' ')[0].rstrip(',;')
    v0 = val.split(' ')[0]
    if re.search(r'\d', tok) and len(tok) >= 5 and (v0.lower() in tok.lower() or tok.lower() in val.lower()):
        if tok.lower() in val.lower():
            return val          # e.g. 'G6K-2F-Y DC4.5'
        return tok
    return ''

# Short descriptions for non-passive parts, keyed by order PN (or Value)
SD = {
    'BAV70': 'Diode/Dual SW 100V/CC', 'SMF24A': 'TVS/24V/200W/Uni', 'BZX84-C13,215': 'Zener/13V/250mW',
    'SMBJ16A': 'TVS/16V/600W/Uni', 'SMBJ22A': 'TVS/22V/600W/Uni', 'BAT54GWJ': 'Schottky/30V/200mA',
    'SMAJ6.0A': 'TVS/6V/400W/Uni', 'ESDS312DBVR': 'ESD Protection/3.6V', 'SMBJ10A': 'TVS/10V/600W/Uni',
    'SM712': 'TVS Array/-7~+12V (RS-485)', 'XGD10603': 'Polymer ESD/32V', 'BAT54W,115': 'Schottky/30V/200mA',
    'SMBJ18A': 'TVS/18V/600W/Uni', 'SMBJ22CA': 'TVS/22V/600W/Bi', 'TVS2200DRVR': 'Flat-clamp TVS/22V',
    'SF-1206F150-2': 'Chip Fuse/1.5A Fast/63V', 'MF-MSMF075/24X-2': 'PTC Fuse/0.75A hold/24V',
    'MF-MSMF200/16X-2': 'PTC Fuse/2A hold/16V', '0454007.MR': 'SMD Fuse/7A Slow/75V',
    'BLM21PG601SN1D': 'BEAD/600ohm/1.4A/2012',
    'B2B-XH-A': 'CONN/2.5mm_2P (Power button)', 'B4B-XH-A': 'CONN/2.5mm_4P (HMI)',
    'B3P-VH(LF)(SN)': 'CONN/3.96mm_3P (Battery)', 'B2P-VH(LF)(SN)': 'CONN/3.96mm_2P (ADP)',
    '1x07 P2.54 vertical male (temporary)': 'HEADER_PIN/7P/DIP', 'C-RSMF.RA4R.1': 'SMA Conn. R/A 50ohm',
    'PJ-063BH': 'DC Jack/2.5mm/8A/24V', 'G6K-2F-Y DC4.5': 'Signal Relay DPDT/4.5V coil',
    'B3S-1000': 'Tact Switch/SMD', 'BSS84LT1G': 'P-MOSFET/-50V/-130mA', 'BSS138': 'N-MOSFET/50V/220mA',
    'NTMFS5C468NLT1G': 'N-MOSFET/40V/10.3mohm', 'MMBT3904': 'NPN BJT/40V/200mA', 'EPC8010': 'eGaN FET/100V/4A',
    'SiS443DN-T1-GE3': 'P-MOSFET/-40V', 'CSD18543Q3A': 'N-MOSFET/60V/9.9mohm',
    'CST2010-100LD': 'Current Transformer 1:100', 'NTCS0603E3103JHT': 'NTC/10k/B3960/0603',
    'LMR33630ADDAR': 'DCDC Buck/36V/3A', 'TLV9061IDBVR': 'OPAMP/1CH/10MHz', 'TLV75733PDBVR': 'LDO/3.3V/1A',
    'TPS26624DRCR': 'eFuse/60V', 'TPS552882RPMR': 'DCDC Buck-Boost/36V', 'LMR36520FADDAR': 'DCDC Buck/65V/2A',
    '74HC595D': 'Shift Register/8bit', 'ULN2803ADWR': 'Darlington Array/8CH', 'TLV3501AIDBVR': 'Comparator/4.5ns',
    'SN74LVC1G17DBVR': 'Schmitt Buffer/1CH', 'XR33194ESBTR': 'RS-422/485 Driver',
    'XR33183ESBTR': 'RS-422/485 Receiver', 'LMG1020YFFR': 'GaN Gate Driver/Low-side',
    'TLV9064IPWR': 'OPAMP/4CH/10MHz', '74AHC1G32GW': 'OR Gate/1CH', '74AHC1G86GW': 'XOR Gate/1CH',
    'SN74LVC74AD': 'Dual D-FF', 'BQ24610RGER': 'Li-ion Charger Ctrl/1-6cell', 'LM74700QDBVRQ1': 'Ideal Diode Ctrl/65V',
    'BoT-nLE523D': 'BLE UART Module', 'NUCLEO-H753ZI': 'MCU Board/STM32H753ZI', 'Piezo': 'Piezo Buzzer/3.3V',
    # MVP-PPG
    'SPHV15-01ETG': 'TVS/15V/Uni', 'ESD351DPYR': 'ESD TVS/3.6V/1.8pF', 'BLM15PX601SN1D': 'BEAD/600ohm/0.9A/1005',
    'MAXM86161EFD+': 'PPG Optical Module/AFE/I2C', 'STM32C011F6U6': 'MCU/Cortex-M0+/32KB Flash',
    'MCP1703A-5002E/MC': 'LDO/5V/250mA', 'MCP1703A-3302E/MC': 'LDO/3.3V/250mA',
}

def pkg_size(kind, size):
    """CAP-/RES- package label; metric 0603 (= inch 0201) is spelled out so it is not read as inch 0603."""
    return f'{kind}-{size}' + ('(0201)' if size == '0603' else '')
NOTE = {
    'BAT54GWJ': 'SOD-123 전용: BAT54(SOT-23)/BAT54W(SOT-323) 주문 금지',
    'BAT54W,115': 'SOT-323: BAT54GWJ(SOD-123)와 혼동 주의',
    'NXH25VB680M10x16': '삼영 NXH 25V/680uF/10x16 (대체: Nichicon UHW1E681MPD, Rubycon 25ZLH680MEFC10X16)',
    'NUCLEO-H753ZI': '완제품 보드 별도 구매 (MB1364)',
    'BoT-nLE523D': '모듈 별도 구매',
    '1x07 P2.54 vertical male (temporary)': '임시 헤더',
    'XGD10603': '주문 형번(접미사) 확인 필요',
    'SM712': '주문 형번(예: SM712-02HTG) 확인 필요',
}

# ---------------------------------------------------------------- classify each reference
items = []
review = []          # (ref, issue)
for r in rows:
    ref, val, fp_full, desc = r['Ref'], r['Value'], r['Footprint'], r['Desc']
    fp = fp_full.split(':')[-1]
    size = size_of(fp)
    attr = fp_attr_of(fp_full)
    if attr is None:
        review.append((ref, f'풋프린트 파일 없음 ({fp_full}) - AUTO로 가정'))
        attr = ''
    level = 'DIP' if 'through_hole' in attr else 'AUTO'
    dnp = r['DNP'] == 'DNP'
    vendor, pn_m, alt = split_vendor(r['Mfr'])
    sup = SUPP.get(ref, {})
    pn = r['MPN'] or r.get('MPN2', '') or pn_m or sup.get('MPNResolved', '')
    vendor = vendor or sup.get('ManufacturerResolved', '')
    notes = []
    if alt:
        notes.append(alt)
    for sn in sup.get('Notes', []):
        mx = re.search(r'예시:?\s*([A-Z0-9][A-Z0-9-]{7,})', sn)
        if mx:
            tail = sn[mx.end():].lstrip('는은 ').strip()     # e.g. 'X7R 적합성 확인 필요'
            notes.append('참고 MPN 예시(미확정): ' + mx.group(1) + (f' ({tail})' if tail else ''))
        if '지급품' in sn:
            notes.append('지급품(consigned) - 기존 BOM 기준')
    pfx = re.match(r'[A-Z]+', ref).group(0)
    it = dict(ref=ref, level=level, dnp=dnp, size=size, pkg=pkg_of(fp), vendor=vendor, pn=pn)

    is_zero_ohm = val.upper() in ('0R', '0') or desc.lower().startswith('0 ohm')
    if pfx == 'C' and 'CP_Radial' in fp:
        it.update(cat='C_DIP', part='680uF/25V', sd='Elec-CAP/NXH/105C')
        sortv = (0,)
    elif pfx == 'C':
        CAP_RE = r'(\d+(?:\.\d+)?)\s*(pF|nF|uF)'
        m = re.search(CAP_RE, desc)
        if m:
            rest = desc[m.end():]
        else:                              # spec lives in Value, e.g. '1uF 6.3V' (MVP-PPG)
            m = re.search(CAP_RE, val)
            rest = val[m.end():] + ' ' + desc if m else ''
        if not m:
            it.update(cat='MLCC', part=val + ' (값 미정)', sd='MLCC/%s' % size, pkg=pkg_size('CAP', size))
            sortv = (0, 0)
            review.append((ref, f'정전용량 미정 (Value "{val}")'))
        else:
            pf = float(m.group(1)) * CAP_MULT[m.group(2)]
            if r['Voltage']:
                v = re.match(r'(\d+)', r['Voltage']).group(1)
            elif r['Rating']:
                v = re.match(r'(\d+)', r['Rating']).group(1)
            else:
                mv = re.search(r'(\d+(?:\.\d+)?)\s*V(?:DC)?\b', rest)
                v = mv.group(1) if mv else ''
            md = re.search(r'X7R|X5R|X7S|C0G|NP0', desc) or re.search(r'X7R|X5R|X7S|C0G|NP0', val)
            diel = (r['Dielectric'] or (md.group(0) if md else '')).replace('NP0', 'C0G')
            tol = r['Tol'].replace('+/-', '') if r['Tol'] else ''
            if sup and (not v or not diel):  # earlier reviewed BOM, e.g. StandardValue '100n/16V/X7R'
                filled = []
                for tok in sup.get('StandardValue', '').split('/')[1:]:
                    if re.fullmatch(r'[\d.]+V', tok) and not v:
                        v = tok[:-1]
                        filled.append(tok)
                    elif re.fullmatch(r'X7R|X5R|X7S|C0G|NP0', tok) and not diel:
                        diel = tok
                        filled.append(tok)
                if filled:
                    notes.append('/'.join(filled) + ': 회로도 미기재 → 기존 BOM 보충값')
                    review.append((ref, '회로도 미기재 사양을 기존 BOM 보충값으로 채움: ' + '/'.join(filled)))
            # Value vs Description cross-check
            mvv = re.search(r'(\d+(?:\.\d+)?)\s*(pF|nF|uF)', val)
            if mvv and abs(float(mvv.group(1)) * CAP_MULT[mvv.group(2)] - pf) > 1e-6:
                review.append((ref, f'Value "{val}" ≠ 설명 정전용량 {cap_fmt(pf)}'))
            mvol = re.search(r'/(\d+)V', val)
            if mvol and v and mvol.group(1) != v:
                review.append((ref, f'Value "{val}" 전압 ≠ 설명 {v}V'))
            if not v:
                review.append((ref, '정격전압 미기재'))
            if not diel:
                review.append((ref, f'유전체 미기재 ("{desc[:40]}")'))
            part, sd = mlcc_strings(pf, v, diel, tol, size)
            it.update(cat='MLCC', part=part, sd=sd, pkg=pkg_size('CAP', size))
            it['_mlcc'] = (pf, v, diel, tol)
            sortv = (pf, float(v or 0))
    elif pfx == 'R' or (is_zero_ohm and pfx in ('F', 'FB')):
        if is_zero_ohm:
            ohm, tol = 0.0, ''
        else:
            m = re.match(r'(\d+(?:\.\d+)?)\s*(mR|R|k|M)?', val)
            mult = {'mR': 1e-3, 'R': 1, None: 1, 'k': 1e3, 'M': 1e6}[m.group(2)]
            ohm = float(m.group(1)) * mult
            mt = re.search(r'(\d+(?:\.\d+)?)\s*%', val) or re.search(r'(\d+(?:\.\d+)?)\s*%', desc)
            tol = mt.group(1) + '%' if mt else '5%'
            ms = re.search(r'/(\d+(?:\.\d+)?)%', sup.get('StandardValue', ''))
            if not mt and ms:
                tol = ms.group(1) + '%'
                notes.append(f'{tol}: 회로도 미기재 → 기존 BOM 보충값')
                review.append((ref, f'회로도 미기재 사양을 기존 BOM 보충값으로 채움: {tol}'))
        if ohm == 0:
            vs = '0'
        elif ohm < 1:
            vs = num(ohm)
        elif ohm < 1e3:
            vs = num(ohm)
        elif ohm < 1e6:
            vs = num(ohm / 1e3) + 'k'
        else:
            vs = num(ohm / 1e6) + 'M'
        mw = re.search(r'(\d+(?:\.\d+)?)\s*W\b', val) or re.search(r'(\d+(?:\.\d+)?)\s*W\b', desc)
        watt = mw.group(1) + 'W' if mw and float(mw.group(1)) >= 0.25 else ''
        extra = []
        if desc.lower().startswith('thin film') or 'ppm' in desc:
            extra.append('25ppm')
        if any(k in desc.lower() for k in ('anti-surge', 'pulse-proof', 'high-pulse')):
            extra.append('Pulse-proof')
            if not pn:
                notes.append('Pulse-proof 품 필수 (예: Vishay CRCW-HP)')
        if 'current-sense' in desc.lower():
            extra.append('Current-sense')
        if pfx in ('F', 'FB'):
            notes.append(('퓨즈' if pfx == 'F' else '비드') + ' 자리 - 0Ω 기본 실장' + (' (≥2A)' if pfx == 'FB' else ''))
        sz = size or '?'
        part = vs + (f'/{tol}' if tol and tol != '5%' else '') + ('' if sz == '1608' else f'/{sz}')
        sd = '/'.join(x for x in ['R', watt, tol if tol and tol != '5%' else '', *extra, sz] if x)
        it.update(cat='Chip-Res', part=part, sd=sd, pkg=pkg_size('RES', sz))
        sortv = (ohm, tol)
    else:
        if not pn:
            pn = pn_from_desc(val, desc)
        key = pn or val
        if pfx == 'TP' or re.search(r'CablePad|TestPad|TestPoint', fp):
            cat = 'PCB Pad'                 # copper only - nothing to buy
        elif pfx == 'F':
            cat = 'Fuse'
        elif pfx == 'FB':
            cat = 'Bead'
        elif pfx == 'L':
            cat = 'L'
        elif pfx == 'D' and 'LED' in fp:
            cat = 'LED'
        elif pfx == 'D' and re.search(r'TVS|ESD', desc):
            cat = 'TVS'
        elif pfx == 'D':
            cat = 'Diode'
        elif pfx == 'J':
            cat = 'Header' if 'PinHeader' in fp else ('Wafer' if 'JST' in fp else 'CONN')
        elif pfx == 'K':
            cat = 'Relay'
        elif pfx == 'Q':
            cat = 'TR' if 'transistor' in desc.lower() and 'MOSFET' not in desc else 'FET'
        elif pfx == 'T':
            cat = 'TRANS'
        elif pfx == 'TH':
            cat = 'Thermistor'
        elif pfx == 'SW':
            cat = 'switch'
        elif pfx == 'BZ':
            cat = 'Buzzer'
        elif pfx == 'A' or 'Module' in SD.get(key, '') or 'module' in desc.lower()[:60]:
            cat = 'Module'
        else:
            cat = 'IC'
        if pfx == 'J' and 'XH' in fp:  # JST XH: order PN from footprint
            key = pn = 'B2B-XH-A' if '1x02' in fp else 'B4B-XH-A'
        sd = SD.get(key, '')
        if cat == 'L':
            mi = re.search(r'Isat\s*([\d.]+)\s*A', desc)
            mt = re.search(r'\+/-\s*(\d+)\s*%', desc)
            sd = f'L/{val}' + (f'/Isat {mi.group(1)}A' if mi else '') + (f'/±{mt.group(1)}%' if mt and not mi else '')
            if 'filter inductor' in desc:
                sd += ' (Filter)'
        if cat == 'LED':
            sd = 'LED/B/1608'
        if cat == 'TRANS' and 'RM6' in fp:
            sd = 'Output Transformer/RM6 PC200/5T+5T:10T'
            notes.append('주문 제작 권선: ' + pn)
            pn = 'Custom (TDK RM6 PC200)'
            vendor = 'Custom wound'
        if cat == 'PCB Pad':
            mdia = re.search(r'D(\d+(?:\.\d+)?)', fp)
            sd = ('Cable' if 'Cable' in fp else 'Test') + ' Pad' + (f' Ø{mdia.group(1)}' if mdia else '') + ' (부품 없음)'
        if key in NOTE:
            notes.append(NOTE[key])
        if not sd:
            review.append((ref, f'Short Description 매핑 없음 ({key})'))
            sd = desc[:40]
        part = val
        if cat == 'Diode' and pn:
            part = pn.split(',')[0]          # 'BAT54' and 'BAT54GWJ' are the same order item
        if cat == 'Header':
            part = 'CON7/2.54mm'
        if cat == 'Wafer':
            part = {'XH-2P': 'XH2.50-2P', 'XH-4P': 'XH2.50-4P', 'VH-2P': 'VH3.96-2P', 'VH-3P': 'VH3.96-3P'}[it['pkg']]
        if cat == 'Bead':
            it['pkg'] = 'L' + size
        it.update(cat=cat, part=part, sd=sd)
        if cat == 'Module' and level == 'AUTO' and pfx == 'A':
            it['level'] = 'DIP'
        sortv = (key,)
    if pn and pn not in (it.get('pn') or ''):
        it['pn'] = pn
    it['pn'] = pn
    if cat_note := (NOTE.get(pn) if it['cat'] in ('C_DIP',) else None):
        notes.append(cat_note)
    if dnp:
        notes.insert(0, 'DNP(미실장)')
    if not it['vendor'] and it['cat'] != 'PCB Pad':
        it['vendor'] = 'Any' if it['cat'] in ('MLCC', 'Chip-Res', 'LED') else ''
        if not it['vendor']:
            review.append((ref, f'제조사 미기재 ({it["part"]})'))
    if it['cat'] not in ('MLCC', 'Chip-Res', 'LED', 'PCB Pad') and not it['pn']:
        review.append((ref, f'주문 형번(MPN) 미기재 ({it["part"]})'))
    if it['cat'] == 'PCB Pad':
        it['level'] = 'PCB'
    elif dnp:
        it['level'] = 'DNP'
    it['notes'] = notes
    it['sortv'] = sortv
    items.append(it)

# ---------------------------------------------------------------- MLCC voltage unification
# Same capacitance / size / dielectric / tolerance / MPN with only the rated voltage differing
# -> order one part at the highest voltage. DNP rows are kept apart from mounted rows (level in key).
def mlcc_key(it):
    pf, v, diel, tol = it['_mlcc']
    return (it['level'], pf, diel, tol, it['size'], it['pn'])

vmax = {}
for it in items:
    if '_mlcc' in it and it['_mlcc'][1]:
        k = mlcc_key(it)
        vmax[k] = max(vmax.get(k, 0), float(it['_mlcc'][1]))
raised = defaultdict(list)          # (cap, size, diel, old V, new V) -> refs
for it in items:
    if '_mlcc' not in it or not it['_mlcc'][1]:
        continue
    pf, v, diel, tol = it['_mlcc']
    top = vmax[mlcc_key(it)]
    if float(v) < top:
        nv = num(top)
        raised[(pf, it['size'], diel, v, nv)].append(it)
        it['part'], it['sd'] = mlcc_strings(pf, nv, diel, tol, it['size'])
        it['sortv'] = (pf, top)
for (pf, size, diel, v, nv), its in raised.items():
    refs = ','.join(sorted((i['ref'] for i in its), key=refkey))
    for i in its:
        i['notes'].append(f'{v}V→{nv}V 상향 통합: {refs}')
    review.append((refs, f'MLCC 전압 통합: {cap_fmt(pf)} {diel} {size} {v}V → {nv}V'))

# ---------------------------------------------------------------- group
CAT_ORDER = ['MLCC', 'C_DIP', 'Chip-Res', 'Fuse', 'Bead', 'L', 'LED', 'Diode', 'TVS',
             'Header', 'Wafer', 'CONN', 'switch', 'Relay', 'TR', 'FET', 'TRANS', 'Thermistor',
             'Buzzer', 'IC', 'Module', 'PCB Pad']
SIZE_ORDER = ['0603', '1005', '1608', '2012', '3216', '3225', '4532', '6332']

groups = OrderedDict()
for it in items:
    k = (it['level'], it['cat'], it['part'], it['pn'], it['sd'], it['pkg'])
    g = groups.setdefault(k, dict(items=[], **{x: it[x] for x in ('level', 'cat', 'part', 'pn', 'sd', 'pkg', 'size', 'sortv')}))
    g['items'].append(it)

def uniq(seq):
    out = []
    for s in seq:
        if s and s not in out:
            out.append(s)
    return out

glist = []
for g in groups.values():
    g['refs'] = sorted((i['ref'] for i in g['items']), key=refkey)
    g['vendor'] = ' / '.join(uniq(i['vendor'] for i in g['items']))
    g['notes'] = '; '.join(uniq(n for i in g['items'] for n in i['notes']))
    glist.append(g)

def gsort(g):
    sec = {'AUTO': 0, 'DIP': 0, 'DNP': 1, 'PCB': 2}[g['level']]
    so = SIZE_ORDER.index(g['size']) if g['size'] in SIZE_ORDER else 99
    if g['cat'] in ('IC', 'Module', 'FET', 'TR', 'TVS', 'Diode', 'Fuse', 'CONN', 'Wafer', 'Header', 'Relay', 'TRANS', 'Thermistor', 'L', 'Bead', 'PCB Pad'):
        return (sec, CAT_ORDER.index(g['cat']), 0, refkey(g['refs'][0]))
    return (sec, CAT_ORDER.index(g['cat']), so, g['sortv'], refkey(g['refs'][0]))

glist.sort(key=gsort)

# sanity: every reference exactly once
all_refs = [ref for g in glist for ref in g['refs']]
assert len(all_refs) == len(set(all_refs)) == len(rows), (len(all_refs), len(rows))

# ---------------------------------------------------------------- workbook
FONT = '맑은 고딕'
f9 = Font(name=FONT, size=9)
f9b = Font(name=FONT, size=9, bold=True)
f9blue = Font(name=FONT, size=9, color='0000FF')
thin = Side(style='thin', color='000000')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)
right = Alignment(horizontal='right', vertical='center')
hdr_fill = PatternFill('solid', fgColor='D9D9D9')
sec_fill = PatternFill('solid', fgColor='F2F2F2')

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'BOM'
ws.sheet_view.showGridLines = False

COLS = [('No.', 5), ('LEVEL', 7), ('Value', 10), ('Part', 20), ('Part number', 22),
        ('Short Description', 30), ('PACKAGE', 17), ("Q'TY", 6), ('LOCATION', 62),
        ('필요수량', 8), ('Vendor', 18), ('기타', 34)]
for i, (_, w) in enumerate(COLS, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# header block (rows 1-6)
ws.merge_cells('A1:F6')
ws['A1'] = 'BILL OF MATERIAL LIST\n_Product List'
ws['A1'].font = Font(name=FONT, size=12, underline='single')
ws['A1'].alignment = center
labels = ['P/L NO. :', 'RETI MODEL :', 'BUYER MODEL :', '작성일자 :', '담당 :', '수량(EA) :']
values = ['', a.model, '', datetime.date.today().isoformat(), '', 5]
for i, (lab, v) in enumerate(zip(labels, values)):
    rr = 1 + i
    ws.cell(rr, 7, lab).font = Font(name=FONT, size=8)
    ws.cell(rr, 7).alignment = left
    ws.merge_cells(start_row=rr, start_column=8, end_row=rr, end_column=9)
    c = ws.cell(rr, 8, v)
    c.font = f9b if lab.startswith('수량') else f9
    c.alignment = left
    if lab.startswith('수량'):
        c.font = Font(name=FONT, size=9, bold=True, color='0000FF')
        c.fill = PatternFill('solid', fgColor='FFFF00')
        c.comment = Comment('생산 수량(EA). 필요수량 = Q\'TY x 이 값. 업로드된 양식 예시값 5를 기본으로 둠 - 실제 수량으로 수정하세요.', 'BOM')
for rr in range(1, 7):
    for cc in range(1, 10):
        ws.cell(rr, cc).border = Border(
            left=thin if cc in (1, 7, 8) else None, right=thin if cc in (6, 9) else None,
            top=thin if rr == 1 or cc >= 7 else None, bottom=thin if rr == 6 or cc >= 7 else None)
QTY_CELL = '$H$6'
for rr in range(1, 7):
    ws.row_dimensions[rr].height = 13

HR = 8
for i, (name, _) in enumerate(COLS, 1):
    c = ws.cell(HR, i, name)
    c.font = f9b
    c.alignment = center
    c.fill = hdr_fill
    c.border = box
ws.row_dimensions[HR].height = 22
ws.freeze_panes = ws.cell(HR + 1, 1)

row = HR + 1
no = 0
first_data = row
section_written = set()
SECTION_TITLE = {'DNP': 'DNP (미실장) — Q\'TY는 참고용, 필요수량 0',
                 'PCB': 'PCB 패드(테스트·케이블 패드) — 부품 없음, 필요수량 0'}

def est_lines(text, width_chars):
    return max(1, math.ceil(len(text) / width_chars)) if text else 1

for g in glist:
    if g['level'] in SECTION_TITLE and g['level'] not in section_written:
        section_written.add(g['level'])
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLS))
        c = ws.cell(row, 1, SECTION_TITLE[g['level']])
        c.font = f9b
        c.fill = sec_fill
        c.alignment = left
        for cc in range(1, len(COLS) + 1):
            ws.cell(row, cc).border = box
        row += 1
    no += 1
    loc = ','.join(g['refs'])
    vals = [no, g['level'], g['cat'], g['part'], g['pn'], g['sd'], g['pkg'], len(g['refs']), loc,
            f'=IF(OR(B{row}="DNP",B{row}="PCB"),0,H{row}*{QTY_CELL})', g['vendor'], g['notes']]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row, i, v)
        c.font = f9
        c.border = box
        c.alignment = left if i in (4, 5, 6, 9, 11, 12) else center
    ws.cell(row, 4).alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    lines = max(est_lines(loc, 70), est_lines(g['notes'], 36), est_lines(g['pn'], 24), est_lines(g['part'], 20))
    ws.row_dimensions[row].height = max(13.5, 12.5 * lines)
    row += 1
last_data = row - 1

# totals
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
ws.cell(row, 1, 'TOTAL (실장 기준, DNP/PCB 제외)').font = f9b
ws.cell(row, 1).alignment = right
ws.cell(row, 8, f'=SUMIFS(H{first_data}:H{last_data},B{first_data}:B{last_data},"<>DNP",B{first_data}:B{last_data},"<>PCB")').font = f9b
ws.cell(row, 9, f'=COUNTIFS(B{first_data}:B{last_data},"<>DNP",B{first_data}:B{last_data},"<>PCB",B{first_data}:B{last_data},"<>")&" 품목"').font = f9b
ws.cell(row, 10, f'=SUM(J{first_data}:J{last_data})').font = f9b
for cc in range(1, len(COLS) + 1):
    ws.cell(row, cc).border = box
    ws.cell(row, cc).fill = hdr_fill
    if cc in (8, 10):
        ws.cell(row, cc).alignment = center
ws.cell(row, 9).alignment = left
total_row = row

# print setup
ws.page_setup.orientation = 'landscape'
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = f'{HR}:{HR}'
ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)
ws.oddFooter.center.text = '&P / &N'
ws.auto_filter.ref = f'A{HR}:{get_column_letter(len(COLS))}{last_data}'

# ---------------------------------------------------------------- review sheet
rv = wb.create_sheet('검토사항')
rv.sheet_view.showGridLines = False
rv.column_dimensions['A'].width = 22
rv.column_dimensions['B'].width = 100
def refs_brief(refs, n=6):
    refs = sorted(refs, key=refkey)
    return ','.join(refs[:n]) + (f' 외 {len(refs) - n}개' if len(refs) > n else '')

excl_txt = '; '.join(f'{k} {len(v)}개({refs_brief(v)})' for k, v in sorted(excluded.items()))
info = [
    ('출처', a.source),
    ('기준 커밋', COMMIT),
    ('추출 방법', 'KiCad 10 kicad-cli sch export bom (참조별) → 값/전압/유전체/공차/사이즈/MPN 기준 그룹화'),
    ('추출일', datetime.date.today().isoformat()),
    ('BOM 포함 참조 수', len(rows)),
    ('실장 참조 수', sum(1 for i in items if i['level'] in ('AUTO', 'DIP'))),
    ('DNP 참조 수', sum(1 for i in items if i['level'] == 'DNP')),
    ('PCB 패드 참조 수', sum(1 for i in items if i['level'] == 'PCB')),
    ('BOM 제외 심볼', (excl_txt + ' — 회로도에서 in_bom=no') if excluded else '없음'),
    ('보충값 출처', a.supp and ('기존 검토 BOM ' + os.path.basename(a.supp) + ' (회로도에 없는 사양만 채움, 기타란에 표시)') or '없음'),
    ('Vendor "Any"', '회로도에 제조사가 없는 범용 MLCC/칩저항/LED. 발주처 지정 필요'),
    ('필요수량', 'BOM 시트 H6(수량 EA) × Q\'TY. DNP/PCB 행은 0'),
]
rv['A1'] = '추출 정보'
rv['A1'].font = Font(name=FONT, size=11, bold=True)
r0 = 2
for k, v in info:
    rv.cell(r0, 1, k).font = f9b
    rv.cell(r0, 2, v).font = f9
    rv.cell(r0, 2).alignment = Alignment(horizontal='left', vertical='center')
    for cc in (1, 2):
        rv.cell(r0, cc).border = box
    r0 += 1
r0 += 1
rv.cell(r0, 1, '확인 필요 항목').font = Font(name=FONT, size=11, bold=True)
r0 += 1
for i, h in enumerate(['Reference', '내용'], 1):
    c = rv.cell(r0, i, h)
    c.font = f9b
    c.fill = hdr_fill
    c.border = box
r0 += 1
# collapse identical issues into one row listing refs
agg = OrderedDict()
for ref, issue in review:
    agg.setdefault(issue, []).append(ref)
for issue, refs in sorted(agg.items(), key=lambda kv: (kv[0].split(' ')[0], refkey(kv[1][0]))):
    rv.cell(r0, 1, ','.join(sorted(set(refs), key=refkey))).font = f9
    rv.cell(r0, 2, issue).font = f9
    for cc in (1, 2):
        rv.cell(r0, cc).border = box
        rv.cell(r0, cc).alignment = Alignment(vertical='center', wrap_text=True)
    r0 += 1

wb.save(OUT)

# ---------------------------------------------------------------- PDF (same layout, static values)
from reportlab.lib.pagesizes import A3, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

pdfmetrics.registerFont(TTFont('Malgun', r'C:\Windows\Fonts\malgun.ttf'))
pdfmetrics.registerFont(TTFont('MalgunB', r'C:\Windows\Fonts\malgunbd.ttf'))
EA = values[-1]
ps = ParagraphStyle('c', fontName='Malgun', fontSize=6.5, leading=8)
psc = ParagraphStyle('cc', parent=ps, alignment=1)
psr = ParagraphStyle('cr', parent=ps, alignment=2)
psb = ParagraphStyle('b', fontName='MalgunB', fontSize=6.5, leading=8, alignment=1)
P = lambda t, s=ps: Paragraph(str(t).replace('&', '&amp;').replace('<', '&lt;'), s)

pw, ph = landscape(A3)
widths_rel = [w for _, w in COLS]
avail = pw - 16 * mm
cw = [avail * w / sum(widths_rel) for w in widths_rel]

head = Table([
    [Paragraph('<u>BILL OF MATERIAL LIST</u><br/><u>_Product List</u>', ParagraphStyle('t', fontName='Malgun', fontSize=11, leading=15, alignment=1)),
     P(labels[0]), P(values[0])],
    ['', P(labels[1]), P(values[1])], ['', P(labels[2]), P(values[2])],
    ['', P(labels[3]), P(values[3])], ['', P(labels[4]), P(values[4])],
    ['', P(labels[5]), Paragraph(f'<b>{EA}</b>', ps)],
], colWidths=[sum(cw[:6]), cw[6], cw[7] + cw[8]], rowHeights=[9] * 6)
head.setStyle(TableStyle([
    ('SPAN', (0, 0), (0, 5)), ('BOX', (0, 0), (0, 5), 0.6, colors.black),
    ('GRID', (1, 0), (2, 5), 0.4, colors.black), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ('TOPPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
]))
head.hAlign = 'LEFT'

data = [[P(n, psb) for n, _ in COLS]]
style = [('GRID', (0, 0), (-1, -1), 0.4, colors.black), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#D9D9D9')),
         ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 1.5),
         ('BOTTOMPADDING', (0, 0), (-1, -1), 1.5), ('LEFTPADDING', (0, 0), (-1, -1), 2), ('RIGHTPADDING', (0, 0), (-1, -1), 2)]
n = 0
sec_done = set()
tot_q = tot_need = 0
for g in glist:
    if g['level'] in SECTION_TITLE and g['level'] not in sec_done:
        sec_done.add(g['level'])
        data.append([Paragraph('<b>' + SECTION_TITLE[g['level']] + '</b>', ps)] + [''] * (len(COLS) - 1))
        ri = len(data) - 1
        style += [('SPAN', (0, ri), (-1, ri)), ('BACKGROUND', (0, ri), (-1, ri), colors.HexColor('#F2F2F2'))]
    n += 1
    q = len(g['refs'])
    need = 0 if g['level'] in ('DNP', 'PCB') else q * EA
    if g['level'] not in ('DNP', 'PCB'):
        tot_q += q
    tot_need += need
    data.append([P(n, psc), P(g['level'], psc), P(g['cat'], psc), P(g['part'], psr), P(g['pn']), P(g['sd']),
                 P(g['pkg'], psc), P(q, psc), P(','.join(g['refs'])), P(need, psr), P(g['vendor'], psc), P(g['notes'])])
data.append([Paragraph('<b>TOTAL (실장 기준, DNP/PCB 제외)</b>', psr)] + [''] * 6 +
            [Paragraph(f'<b>{tot_q}</b>', psc), Paragraph(f'<b>{n - sum(1 for g in glist if g["level"] in ("DNP", "PCB"))} 품목</b>', ps),
             Paragraph(f'<b>{tot_need}</b>', psr), '', ''])
ri = len(data) - 1
style += [('SPAN', (0, ri), (6, ri)), ('BACKGROUND', (0, ri), (-1, ri), colors.HexColor('#D9D9D9'))]
tbl = Table(data, colWidths=cw, repeatRows=1)
tbl.setStyle(TableStyle(style))

def footer(c, d):
    c.setFont('Malgun', 7)
    c.drawCentredString(pw / 2, 6 * mm, f'{d.page}')
    c.drawRightString(pw - 8 * mm, 6 * mm, f'{a.model} BOM · {COMMIT}')

doc = SimpleDocTemplate(OUT[:-5] + '.pdf', pagesize=landscape(A3), leftMargin=8 * mm, rightMargin=8 * mm,
                        topMargin=8 * mm, bottomMargin=10 * mm, title=f'{a.model} BOM')
doc.build([head, Spacer(1, 4 * mm), tbl], onFirstPage=footer, onLaterPages=footer)

with open(OUT + '.txt', 'w', encoding='utf-8') as fh:
    n = 0
    for g in glist:
        n += 1
        fh.write(' | '.join(str(x) for x in [n, g['level'], g['cat'], g['part'], g['pn'], g['sd'], g['pkg'],
                                             len(g['refs']), ','.join(g['refs']), g['vendor'], g['notes'][:80]]) + '\n')
print('groups', len(glist), 'refs', len(all_refs), 'review issues', len(agg), 'rows', HR + 1, '-', last_data, 'total row', total_row)
for issue, refs in agg.items():
    print('  ', ','.join(refs[:8]) + ('…' if len(refs) > 8 else ''), '|', issue)
