#!/usr/bin/env python3
"""Export the current schematic BOM as grouped CSV/XLSX, without DNP parts.

Run from any directory: python /path/to/MVP-PPG/kicad/bom_v08.py
Requires KiCad 10's kicad-cli and openpyxl. Set KICAD_CLI if it is not on PATH.
Only the requested output directory is written; importing this module does no IO.
"""
from argparse import ArgumentParser
from collections import defaultdict
import csv
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET

REPO = Path(__file__).resolve().parent.parent
STEM = "MVP-PPG_BOM_부품별통합_DNP제외_값표준화_v2"
HEADERS = ["No.", "레퍼런스", "수량", "값", "MPN", "제조사", "패키지", "참고 데이터시트", "비고"]
PACKAGE = {
    "SleepBud:MCP1703A_MC_DFN8_3x2mm_P0.5mm": "DFN-8 MC, 3 × 2 mm, P0.5 mm (노출패드)",
    "SleepBud:ST_UFQFPN-20_3x3mm_P0.5mm_NoSilk": "UFQFPN-20, 3 × 3 mm, P0.5 mm",
    "SleepBud:MAXM86161_OLGA-14": "OLGA-14, 2.9 × 4.3 mm",
    "Package_DFN_QFN:ST_UFQFPN-20_3x3mm_P0.5mm": "UFQFPN-20, 3 × 3 mm, P0.5 mm",
    "Package_TO_SOT_SMD:SOT-23-5": "SOT-23-5",
    "SleepBud:TSOT23-6_MaxLinear": "TSOT-23-6",
    "Package_TO_SOT_SMD:SOT-23": "SOT-23",
    "Package_DFN_QFN:DFN-8-1EP_2x3mm_P0.5mm_EP0.61x2.2mm": "DFN-8, 2 × 3 mm, P0.5 mm (노출패드)",
    "Capacitor_SMD:C_0402_1005Metric": "0402 (1005)",
    "Capacitor_SMD:C_0201_0603Metric": "0201 (0603)",
    "Resistor_SMD:R_0201_0603Metric": "0201 (0603)",
    "Resistor_SMD:R_0402_1005Metric": "0402 (1005)",
    "Inductor_SMD:L_0402_1005Metric": "0402 (1005)",
    "SleepBud:TVS_SOD882_1x0.6": "SOD-882, 1.0 × 0.6 mm",
    "SleepBud:TI_DPY0002A_1x0.6": "DFN1006-2 (DPY), 1.0 × 0.6 mm",
    "SleepBud:CablePad_D1.2": "구리 납땜 패드 Ø1.2 mm",
    "SleepBud:CablePad_D0.8": "구리 테스트패드 Ø0.8 mm",
}
MANUFACTURER = {
    "MAXM86161EFD+": "Analog Devices (Maxim)", "STM32C011F6U6": "STMicroelectronics",
    "SN65LVDS1DBVR": "Texas Instruments", "SN65LVDT2DBVR": "Texas Instruments",
    "MCP1703A-5002E/MC": "Microchip", "MCP1703A-3302E/MC": "Microchip",
    "SPHV15-01ETG": "Littelfuse", "ESD351DPYR": "Texas Instruments", "BLM15PX601SN1D": "Murata",
    "XR33194ESBTR": "MaxLinear", "XR33183ESBTR": "MaxLinear",
    "CDSOT23-SM712": "Bourns",
    "CRCW0402120RFKEDHP": "Vishay", "CRCW040222R0JNEDHP": "Vishay",
}
# Specifications explicitly recorded in the original v0.8 BOM and carried into
# the user-reviewed 2026-10-01 workbook. Match value AND package: a reused ref
# must not silently inherit these old ratings. Example MPNs remain notes only.
SUPPLEMENTS = {
    ("R", "4.7k", "Resistor_SMD:R_0201_0603Metric"): (
        "4.7k/1%", "기존 v0.8 BOM의 1%", "MPN 미확정"),
    ("C", "10uF 6.3V X5R", "Capacitor_SMD:C_0402_1005Metric"): (
        "10u/6.3V/X5R", "현행 회로도", "MPN 미확정. 기존 BOM 예시: GRM155R60J106ME47"),
    ("C", "1uF 6.3V", "Capacitor_SMD:C_0201_0603Metric"): (
        "1u/6.3V/X5R", "기존 v0.8 BOM의 X5R", "MPN 미확정. 기존 BOM 예시: GRM033R60J105MEA2"),
    ("C", "100nF", "Capacitor_SMD:C_0201_0603Metric"): (
        "100n/16V/X7R", "기존 v0.8 BOM의 16V/X7R", "MPN 미확정. 기존 BOM 예시: GRM033R71C104KE14"),
    ("C", "1uF 25V X7R", "Capacitor_SMD:C_0402_1005Metric"): (
        "1u/25V/X7R", "현행 회로도", "MPN 미확정. 기존 예시 GRM155R61E105KA12는 X7R 적합성 확인 필요"),
}
SUPPLEMENT_REFS = {
    "4.7k": {"R1", "R2"}, "10uF 6.3V X5R": {"C1"},
    "1uF 6.3V": {"C4", "C5", "C12", "C14"},
    "100nF": {"C6", "C7", "C11"}, "1uF 25V X7R": {"C13"},
}
LDO_DATASHEET = "${KIPRJMOD}/../datasheet/pdf/MCP1703A.pdf"


def ref_key(ref):
    return tuple(int(s) if s.isdigit() else s for s in re.split(r"(\d+)", ref))


def compact_value(value, family):
    """Keep the electrical rating; omit only the requested unit/tolerance text."""
    value = value.strip().replace("µ", "u").replace("μ", "u")
    if family not in {"R", "C", "L"}:
        return value
    value = re.sub(r"(?<=\d)\s+(?=[pnumkM]?(?:Ω|[Rr]\b|[Oo]hms?\b|F\b|H\b|V\b|W\b|A\b))", "", value)
    value = re.sub(r"(?<=\d)\s+(?=[pnumkM]\b)", "", value)
    value = re.sub(r"\s*/\s*|\s+", "/", value)
    if family == "R":
        value = re.sub(r"(?<=\d)([mkM]?)(?:Ω|[Rr]|[Oo]hms?)", r"\1", value)
        value = re.sub(r"/(?:±)?5%($|/)", r"\1", value)
    elif family == "C":
        value = re.sub(r"(?<=[0-9pnuµm])F\b", "", value)
    else:
        value = re.sub(r"(?<=[0-9pnuµm])H\b", "", value)
    return value.replace("±", "")


def datasheet_filename(source):
    name = unquote(urlsplit(source.replace("\\", "/")).path).rsplit("/", 1)[-1]
    return name if name.lower().endswith(".pdf") else ""


def cli_path():
    configured = os.environ.get("KICAD_CLI")
    if configured:
        return configured
    found = shutil.which("kicad-cli") or shutil.which("kicad-cli.exe")
    if found:
        return found
    for name in ("kicad-cli", "kicad-cli.exe"):
        sibling = Path(sys.executable).with_name(name)
        if sibling.is_file():
            return str(sibling)
    raise RuntimeError("KiCad 10 kicad-cli is required; set PATH or KICAD_CLI")


def export_schematic(schematic):
    with tempfile.TemporaryDirectory(prefix="ppg-bom-") as work:
        netlist = Path(work) / "current.xml"
        result = subprocess.run(
            [cli_path(), "sch", "export", "netlist", "--format", "kicadxml",
             "--output", str(netlist), str(schematic)],
            capture_output=True, text=True, check=False,
        )
        if result.returncode:
            raise RuntimeError(f"KiCad netlist export failed ({result.returncode}):\n{result.stdout}{result.stderr}")
        return ET.parse(netlist).getroot()


def read_parts(xml):
    parts, excluded = [], []
    seen = set()
    for comp in xml.findall("components/comp"):
        ref = comp.attrib["ref"]
        if ref in seen:
            raise ValueError(f"Duplicate schematic reference: {ref}")
        seen.add(ref)
        properties = {p.attrib["name"] for p in comp.findall("property")}
        if properties & {"dnp", "exclude_from_bom"} or ref.startswith("#"):
            excluded.append({"Reference": ref, "DNP": "dnp" in properties,
                             "ExcludeFromBOM": "exclude_from_bom" in properties})
            continue
        fields = {f.attrib["name"]: f.text or "" for f in comp.findall("fields/field")}
        parts.append({
            "Reference": ref, "Value": comp.findtext("value", ""),
            "Footprint": comp.findtext("footprint", ""), "MPN": fields.get("MPN", ""),
            "Datasheet": fields.get("Datasheet", comp.findtext("datasheet", "")),
            "Description": fields.get("Description", comp.findtext("description", "")),
            "Manufacturer": fields.get("Manufacturer", fields.get("제조사", "")),
        })
    if not parts:
        raise ValueError("The schematic contains no included BOM components")
    return sorted(parts, key=lambda p: ref_key(p["Reference"])), excluded


def group_parts(parts):
    groups = defaultdict(list)
    audit = []
    for part in parts:
        ref, original, footprint = part["Reference"], part["Value"], part["Footprint"]
        family = re.match(r"[A-Za-z]+", ref).group()
        if footprint not in PACKAGE:
            raise ValueError(f"{ref}: review package display for {footprint!r} before export")
        value = compact_value(original, family)
        mpn, source = part["MPN"], part["Datasheet"]
        manufacturer = part["Manufacturer"] or MANUFACTURER.get(mpn or original, "")
        notes, origin = [], ["현행 회로도"]
        supplement = SUPPLEMENTS.get((family, original, footprint))
        # Explicit ordering numbers/updated part fields take precedence over
        # generic passive recommendations from the old BOM.
        if supplement and not mpn and ref in SUPPLEMENT_REFS.get(original, set()):
            value, evidence, note = supplement
            if evidence != "현행 회로도":
                origin.append(evidence)
                notes.append(evidence)
            notes.append(note)
        if family == "U" and original in MANUFACTURER:
            mpn = mpn or original
            if original == "MAXM86161EFD+":
                notes.append("기존 BOM: 지급품(consigned)")
            if original.startswith("MCP1703A-") and not source:
                source = LDO_DATASHEET
                notes.append("데이터시트: 저장소의 Microchip DS20005122C 원본")
        if mpn in {"CRCW0402120RFKEDHP", "CRCW040222R0JNEDHP"}:
            notes.append("HP 펄스 대응 저항, 0.2W (P70). 실제 PCB 열·펄스·ESD 성능 검증 필요")
            notes.append("주문형번은 제조사 코드 규칙 기준; 재고·수급 확인 전")
        if mpn == "CDSOT23-SM712":
            notes.append("RS-422 2선 보호, -7/+12V 동작 범위. 실제 PCB ESD 성능 검증 필요")
        if family == "R" and not mpn and not supplement:
            notes.append("MPN 미확정")
        if family == "FB" and (mpn or original) == "BLM15PX601SN1D":
            value = "600Ω@100MHz/0.9A"
            notes.append("정격 0.9A @85°C, DCR 최대 0.23Ω")
            origin.append("현행 회로도 Description의 임피던스/전류")
        if footprint.startswith("SleepBud:CablePad_"):
            value = "케이블 납땜 패드" if footprint.endswith("D1.2") else "SWD 테스트패드"
            notes.append("구리 패드 — 구매 부품 아님")
        filename = datasheet_filename(source) if source else ""
        if not filename and not family == "J":
            notes.append("PDF 미연결(제조사 제품 페이지)" if source else "데이터시트 미연결")
        record = {**part, "StandardValue": value, "MPNResolved": mpn,
                  "ManufacturerResolved": manufacturer, "Package": PACKAGE[footprint],
                  "DatasheetFilename": filename, "DatasheetSource": source,
                  "Source": origin, "Notes": notes}
        # Full source URL participates so distinct datasheets sharing a basename
        # cannot accidentally merge different specifications.
        key = (family, value, mpn, manufacturer, footprint, filename, source, tuple(notes))
        groups[key].append(record)
        audit.append(record)
    rows = []
    for key, records in groups.items():
        family, value, mpn, manufacturer, footprint, filename, source, notes = key
        refs = sorted((r["Reference"] for r in records), key=ref_key)
        notes = list(notes)
        if family == "J":
            notes.append("; ".join(r["Reference"] + "=" + r["Value"] for r in records))
        rows.append([len(rows) + 1, ", ".join(refs), len(refs), value, mpn,
                     manufacturer, PACKAGE[footprint], filename, "\n".join(notes)])
    return rows, audit


def write_workbook(path, rows, excluded):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "BOM_부품별"
    total = sum(r[2] for r in rows)
    dnp = sum(p["DNP"] for p in excluded)
    excluded_bom = sum(p["ExcludeFromBOM"] for p in excluded)
    summaries = [
        "MVP-PPG BOM — 부품별 통합 · DNP 제외 · 값 표준 v2",
        f"현행 회로도 기준 | {len(rows)}그룹 · {total}개 레퍼런스(구리 패드 포함) | DNP {dnp}개 · BOM 제외 {excluded_bom}개",
        "저항 Ω·커패시터 F·인덕터 H 생략. 저항 기본 5% 생략; 명시된 1%·0.1%는 /오차율로 표시. 패키지·데이터시트 파일명 표기.",
        "예시 MPN은 비고에만 기록하며 PDF 원문은 포함하지 않습니다. 회로도 BOM이며 PCB 제조 검증 결과는 별도 확인합니다.",
    ]
    for number, text in enumerate(summaries, 1):
        ws.merge_cells(start_row=number, start_column=1, end_row=number, end_column=9)
        ws.cell(number, 1, text)
        ws.cell(number, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(number, 1).font = Font(name="맑은 고딕", size=14 if number == 1 else 10,
                                     bold=number == 1, color="FFFFFF" if number == 1 else "172B4D")
        if number == 1:
            ws.cell(number, 1).fill = PatternFill("solid", fgColor="17365D")
        ws.row_dimensions[number].height = 32
    ws.append(HEADERS)
    widths = [8, 36, 9, 34, 36, 28, 47, 58, 71]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in rows:
        ws.append(row)
        lines = []
        for column, value in enumerate(row, 1):
            cell = ws.cell(ws.max_row, column)
            # Treat all source text literally; never interpret MPNs/notes as formulas.
            if isinstance(value, str):
                cell.data_type = "s"
            cell.font = Font(name="맑은 고딕", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top",
                                       horizontal="center" if column in (1, 3) else "left")
            lines.append(sum(max(1, math.ceil(sum(2 if ord(c) > 127 else 1 for c in line) /
                                              (widths[column - 1] - 3)))
                             for line in str(value).split("\n")))
        ws.row_dimensions[ws.max_row].height = max(38, 16 * max(lines) + 8)
    table = Table(displayName="MVP_PPG_BOM_Compact", ref=f"A5:I{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "C6"
    ws.print_area = f"A1:I{ws.max_row}"
    ws.print_title_rows = "1:5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    wb.properties.title = summaries[0]
    wb.properties.subject = "현행 회로도; 동일 부품 레퍼런스 통합; 패키지 표기; 데이터시트 파일명"
    wb.save(path)


def generate(schematic, output):
    xml = export_schematic(schematic)
    parts, excluded = read_parts(xml)
    rows, audit = group_parts(parts)
    refs = [ref for row in rows for ref in row[1].split(", ")]
    if len(refs) != len(set(refs)) or set(refs) != {p["Reference"] for p in parts}:
        raise ValueError("BOM grouping changed the reference set")
    output.mkdir(parents=True, exist_ok=True)
    write_workbook(output / f"{STEM}.xlsx", rows, excluded)
    with (output / "ppg_pcb_v08_BOM.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(HEADERS)
        writer.writerows(rows)
    source = {
        "schematic": schematic.name, "export_tool": xml.findtext("design/tool", ""),
        "groups": len(rows), "included_references": len(parts), "excluded": excluded,
        "parts": audit,
    }
    (output / "bom_sources.json").write_text(json.dumps(source, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return rows, excluded


def main(argv=None):
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("--schematic", type=Path, default=REPO / "kicad/ppg_pcb_v08.kicad_sch")
    parser.add_argument("--output-dir", type=Path, default=REPO / "bom")
    args = parser.parse_args(argv)
    rows, excluded = generate(args.schematic.resolve(), args.output_dir.resolve())
    print(f"BOM: {len(rows)} groups, {sum(r[2] for r in rows)} references; "
          f"DNP {sum(p['DNP'] for p in excluded)}, excluded from BOM {sum(p['ExcludeFromBOM'] for p in excluded)}")
    print(args.output_dir.resolve() / f"{STEM}.xlsx")


if __name__ == "__main__":
    main()
